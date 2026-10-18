"""ONE-OFF BACKFILL from your Excel workbook (history before the PDFs take over), plus first-time creation of lookups.json.
Usage: python scripts/build_data.py path/to/LTA_Tracking.xlsx [--before 2026-01]
Months from --before onward are NOT imported: those come from the official LTA PDFs.
"""
import json, re, sys
from pathlib import Path
from openpyxl import load_workbook
sys.path.insert(0, str(Path(__file__).resolve().parent))
import store

SRC = Path(sys.argv[1])
BEFORE = sys.argv[sys.argv.index("--before") + 1] if "--before" in sys.argv else "2026-01"
RAW_SHEETS = ["PC Raw File 2024", "PC Raw File 2025", "PC Raw File 2026"]
BODY = ["HB", "SDN", "MPV", "STW", "SUV", "CPE"]
ID_COLS = {"Make": "make", "Importer Type": "importer", "Fuel Type": "fuel", "Country": "country", "Distibutor": "distributor", "Prem / Lux": "tier"}
MAKE_ALIASES = {"B.M.W.": "BMW", "M.G.": "MG"}
FUEL_ALIASES = {"Petrol-Electric (Plug-In)": "Petrol-Electric (Plug-in)"}
MONTH_RE = re.compile(r"^\d{4}-\d{2}$")
units = lambda v: int(v) if isinstance(v, (int, float)) else 0   # blanks, non-breaking spaces, stray characters -> 0

wb = load_workbook(SRC, read_only=True, data_only=True)
rows, lookup = [], {"make": {}, "distributor": {}}
for name in RAW_SHEETS:
    it = wb[name].iter_rows(values_only=True)
    h1, h2 = next(it), next(it)
    month_cols = {i: v for i, v in enumerate(h1) if isinstance(v, str) and MONTH_RE.match(v)}
    id_cols = {i: ID_COLS[v] for i, v in enumerate(h2) if v in ID_COLS}
    for r in it:
        if not r or not r[0] or str(r[0]).strip().upper() in ("TOTAL", "GRAND TOTAL"):
            continue
        meta = {k: ("" if r[i] is None else str(r[i]).strip()) for i, k in id_cols.items()}
        meta["make"] = MAKE_ALIASES.get(meta["make"], meta["make"])
        meta["fuel"] = FUEL_ALIASES.get(meta["fuel"], meta["fuel"])
        if meta.get("country"):
            lookup["make"][meta["make"]] = {"country": meta["country"], "tier": meta.get("tier", "")}
        if meta.get("distributor"):
            lookup["distributor"][f'{meta["make"]}|{meta["importer"]}'] = meta["distributor"]
        for c, month in month_cols.items():
            if month >= BEFORE:
                continue
            for j, body in enumerate(BODY):
                u = units(r[c + j]) if c + j < len(r) else 0
                if u:
                    rows.append({"m": month, "make": meta["make"], "importer": meta["importer"], "fuel": meta["fuel"], "body": body, "u": u})

if not store.LOOKUPS.exists():
    store.LOOKUPS.write_text(json.dumps(lookup, indent=1, sort_keys=True), encoding="utf-8")
    print("created lookups.json with", len(lookup["make"]), "makes")
print("backfilled months:", store.save_months(rows, overwrite=False))
store.rebuild()
